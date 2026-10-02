"""Bounded local evidence parsers. Content is never evaluated as instructions."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import os
import re
import statistics
import uuid
import zipfile
from pathlib import Path
from xml.etree import ElementTree

MAX_BYTES = 25 * 1024 * 1024
MAX_ROWS = 100000
MAX_COLUMNS = 256
MAX_TEXT = 200000
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif"}


def checkpoint(cancel=None):
    if cancel is not None and cancel.is_set():
        raise InterruptedError("Research cancelled; completed records are retained.")


def _safe_xml(data):
    if b"<!DOCTYPE" in data.upper() or b"<!ENTITY" in data.upper():
        raise ValueError("XML entities/DTDs are not supported.")
    return ElementTree.fromstring(data)


def profile_table(headers, rows, cancel=None):
    headers = [str(v or "").strip() for v in headers]
    if not headers or len(headers) > MAX_COLUMNS or any(not h for h in headers):
        raise ValueError("Dataset needs 1–256 non-empty column names.")
    if len(set(headers)) != len(headers):
        raise ValueError("Duplicate column names make dataset references ambiguous.")
    values = [[] for _ in headers]
    seen, duplicates, count = set(), 0, 0
    preview = []
    for row in rows:
        checkpoint(cancel)
        count += 1
        if count > MAX_ROWS:
            raise ValueError("Dataset exceeds the 100,000-row bounded analysis limit.")
        if len(row) != len(headers):
            raise ValueError(f"Row {count + 1} has {len(row)} values, expected {len(headers)}.")
        cleaned = tuple("" if v is None else str(v).strip() for v in row)
        digest = hashlib.sha256(json.dumps(cleaned).encode()).digest()
        duplicates += digest in seen
        seen.add(digest)
        if count <= 8:
            preview.append(cleaned)
        for col, value in zip(values, cleaned):
            col.append(value)
    columns = []
    warnings = []
    for header, cells in zip(headers, values):
        checkpoint(cancel)
        present = [v for v in cells if v.casefold() not in {"", "na", "n/a", "null", "none"}]
        numeric = []
        invalid = 0
        for value in present:
            try:
                number = float(value)
                if not math.isfinite(number):
                    invalid += 1
                else:
                    numeric.append(number)
            except ValueError:
                pass
        unit_match = re.search(r"\[([^\]]+)\]|\(([^)]+)\)", header)
        col = {"name": header, "missing": count - len(present),
               "unit_label": next((s for s in unit_match.groups() if s), "") if unit_match else "",
               "type": "numeric" if present and len(numeric) + invalid == len(present) else "text",
               "nonfinite": invalid}
        if col["type"] == "numeric" and numeric:
            ordered = sorted(numeric)
            n = len(numeric)
            q1, q3 = (statistics.quantiles(ordered, n=4, method="inclusive")[i] for i in (0, 2)) if n >= 4 else (ordered[0], ordered[-1])
            iqr = q3 - q1
            col.update({"count": n, "mean": statistics.fmean(numeric), "median": statistics.median(numeric),
                        "min": min(numeric), "max": max(numeric),
                        "sample_std": statistics.stdev(numeric) if n > 1 else None,
                        "iqr_outliers": sum(x < q1 - 1.5 * iqr or x > q3 + 1.5 * iqr for x in numeric) if n >= 4 else None,
                        "plot_sample": [[i + 1, float(v)] for i, v in enumerate(cells[:500]) if _finite(v)]})
        elif present:
            col["distinct_values"] = len(set(present))
            col["examples"] = list(dict.fromkeys(present))[:6]
            if numeric:
                warnings.append(f"{header}: mixed numeric and text values; not coerced into a numeric result.")
        columns.append(col)
    if count < 5:
        warnings.append("Very small sample; descriptive summaries do not establish a hypothesis.")
    if duplicates:
        warnings.append(f"{duplicates} duplicate rows; retained, not silently removed.")
    if any(c["missing"] for c in columns):
        warnings.append("Missing values excluded per column; counts may differ.")
    if any(c["nonfinite"] for c in columns):
        warnings.append("Non-finite values excluded and flagged.")
    return {"rows": count, "columns": columns, "duplicate_rows": duplicates, "preview": preview,
            "warnings": warnings, "method": "Descriptive statistics; Tukey 1.5×IQR flags, no automatic outlier deletion",
            "limitations": "Column labels are not verified units; groups, causality and independent variables require study context."}


def _finite(value):
    try:
        return math.isfinite(float(value))
    except (ValueError, TypeError):
        return False


def document_sections(text):
    headings = list(re.finditer(r"(?im)^\s*(?:\d+[. ]+)?(abstract|introduction|methods|materials and methods|results|discussion|limitations|conclusions?|references)\s*[:.]?\s*$", text))
    return {m.group(1).lower(): text[m.end():headings[i + 1].start() if i + 1 < len(headings) else len(text)][:12000]
            for i, m in enumerate(headings)}


class ResearchIngestion:
    def __init__(self, blob_directory: str | Path):
        self.blobs = Path(blob_directory)

    def ingest(self, source: str | Path, cancel=None):
        checkpoint(cancel)
        source = Path(source).resolve(strict=True)
        if not source.is_file() or source.stat().st_size > MAX_BYTES:
            raise ValueError("Research files must be regular files no larger than 25 MiB.")
        with source.open("rb") as stream:
            raw = stream.read(MAX_BYTES + 1)
        if len(raw) > MAX_BYTES:
            raise ValueError("File grew beyond the ingestion size limit.")
        digest = hashlib.sha256(raw).hexdigest()
        ext = source.suffix.lower()
        result = {"name": source.name, "sha256": digest, "bytes": len(raw),
                  "format": ext, "trust": "untrusted_source_data", "warnings": [], "metadata": {}}
        # Parse the exact bytes later stored under their digest, not a second
        # read of a file which may have changed during an experiment.
        if ext in {".csv", ".tsv"}:
            reader = csv.reader(io.StringIO(raw.decode("utf-8-sig")), delimiter="\t" if ext == ".tsv" else ",")
            result["dataset"] = profile_table(next(reader, []), reader, cancel)
        elif ext == ".xlsx":
            _validate_zip(raw)
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
            try:
                tables = []
                for sheet in workbook.worksheets[:20]:
                    checkpoint(cancel)
                    rows = sheet.iter_rows(values_only=True)
                    tables.append({"sheet": sheet.title, **profile_table(next(rows, ()), rows, cancel)})
                result["tables"] = tables
                result["warnings"].append("Excel formulas are source text and are never executed; only the first 20 sheets are inspected.")
            finally:
                workbook.close()
        elif ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(raw))
            if reader.is_encrypted:
                raise ValueError("Encrypted PDF requires an unlocked copy.")
            if len(reader.pages) > 300:
                raise ValueError("PDF exceeds the 300-page limit; split it before importing.")
            pages, length = [], 0
            for number, page in enumerate(reader.pages, 1):
                checkpoint(cancel)
                text = (page.extract_text() or "")[:MAX_TEXT - length]
                pages.append({"page": number, "text": text})
                length += len(text)
                if length >= MAX_TEXT:
                    result["warnings"].append("PDF text truncated at 200,000 characters.")
                    break
            result["pages"] = pages
            result["text"] = "\n".join(p["text"] for p in pages)
            metadata = reader.metadata
            result["metadata"] = {"title": str(getattr(metadata, "title", "") or ""),
                                  "author": str(getattr(metadata, "author", "") or ""),
                                  "page_count": len(reader.pages)}
            result["sections"] = document_sections(result["text"])
            result["warnings"].append("Text extraction only: figures, equations and table layout require verification; scanned pages need OCR.")
        elif ext == ".docx":
            _validate_zip(raw)
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                root = _safe_xml(archive.read("word/document.xml"))
            result["text"] = "\n".join("".join(p.itertext()) for p in root.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"))[:MAX_TEXT]
            result["sections"] = document_sections(result["text"])
        elif ext in IMAGE_EXTENSIONS:
            from PIL import Image
            with Image.open(io.BytesIO(raw)) as image:
                if image.width * image.height > 40000000:
                    raise ValueError("Image exceeds 40 megapixels.")
                image.verify()
                result["metadata"] = {"width": image.width, "height": image.height, "mode": image.mode}
            result["warnings"].append("Image integrity checked; no scientific visual conclusion without a configured vision provider.")
        elif ext in {".json", ".ipynb"}:
            parsed = json.loads(raw.decode("utf-8-sig"))
            result["text"] = json.dumps(parsed, ensure_ascii=False)[:MAX_TEXT]
            result["metadata"]["structure"] = type(parsed).__name__
        elif ext == ".xml":
            root = _safe_xml(raw)
            result["text"] = " ".join(root.itertext())[:MAX_TEXT]
            result["metadata"]["root"] = root.tag
        elif ext in {".txt", ".md", ".log", ".py", ".r", ".js", ".ts", ".c", ".cpp", ".yaml", ".yml"}:
            result["text"] = raw.decode("utf-8-sig")[:MAX_TEXT]
            result["sections"] = document_sections(result["text"])
        else:
            raise ValueError(f"No verified research parser for {ext or 'this file type'}. The file was not analyzed.")
        checkpoint(cancel)
        self.blobs.mkdir(parents=True, exist_ok=True)
        path = self.blobs / digest
        if not path.exists():
            temporary = self.blobs / (digest + "." + uuid.uuid4().hex + ".tmp")
            try:
                temporary.write_bytes(raw)
                os.replace(temporary, path)
            finally:
                temporary.unlink(missing_ok=True)
        elif hashlib.sha256(path.read_bytes()).hexdigest() != digest:
            raise ValueError("Stored evidence failed its checksum; refusing to reuse it.")
        result["blob"] = str(path)
        return result


def _validate_zip(raw):
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        members = archive.infolist()
        if len(members) > 10000 or sum(m.file_size for m in members) > 100 * 1024 * 1024:
            raise ValueError("Document archive exceeds the safe expanded-size limit.")
